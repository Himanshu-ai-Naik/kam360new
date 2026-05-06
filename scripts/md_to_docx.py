#!/usr/bin/env python3
"""Convert AGENT_ROLES_AND_RESPONSIBILITIES.md to .docx with headings, tables, bullets."""
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

def add_paragraph_with_bold(doc, text):
    """Add a paragraph, rendering **bold** as bold runs."""
    if not text.strip():
        return
    p = doc.add_paragraph()
    parts = re.split(r'(\*\*[^*]+\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2])
            run.bold = True
        else:
            p.add_run(part)

def parse_table(lines, start_idx):
    """Parse markdown table from lines starting at start_idx. Returns (rows, next_idx)."""
    rows = []
    i = start_idx
    while i < len(lines):
        line = lines[i]
        if not line.strip().startswith('|'):
            break
        cells = [c.strip() for c in line.split('|')[1:-1]]
        if cells and not all(re.match(r'^[-:]+$', c) for c in cells):  # skip separator row
            rows.append(cells)
        i += 1
    return rows, i

def main():
    root = Path(__file__).resolve().parent.parent
    md_path = root / 'AGENT_ROLES_AND_RESPONSIBILITIES.md'
    out_path = root / 'AGENT_ROLES_AND_RESPONSIBILITIES.docx'

    text = md_path.read_text(encoding='utf-8')
    lines = text.split('\n')
    doc = Document()

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith('# '):
            doc.add_heading(stripped[2:].strip(), level=0)
            i += 1
            continue
        if stripped.startswith('## '):
            doc.add_heading(stripped[3:].strip(), level=1)
            i += 1
            continue
        if stripped.startswith('### '):
            doc.add_heading(stripped[4:].strip(), level=2)
            i += 1
            continue
        if stripped == '---':
            doc.add_paragraph()
            i += 1
            continue
        if stripped.startswith('|') and '|' in stripped:
            rows, next_i = parse_table(lines, i)
            if rows:
                table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                table.style = 'Table Grid'
                for ri, row in enumerate(rows):
                    for ci, cell in enumerate(row):
                        if ci < len(table.rows[ri].cells):
                            table.rows[ri].cells[ci].text = cell
                doc.add_paragraph()
            i = next_i
            continue
        if stripped.startswith('- ') or stripped.startswith('* '):
            bullet = stripped[2:].strip()
            add_paragraph_with_bold(doc, bullet)
            p = doc.paragraphs[-1]
            p.paragraph_format.left_indent = Pt(18)
            p.style = 'List Bullet'
            i += 1
            continue
        if stripped:
            add_paragraph_with_bold(doc, stripped)
            i += 1
            continue
        i += 1

    doc.save(out_path)
    print(f"Wrote {out_path}")

if __name__ == '__main__':
    main()
